from io import BytesIO, StringIO
from collections import defaultdict
from datetime import datetime
import json
import pandas as pd
from app.models.national_material import NationalMaterial
from app.models.material_mapping import MaterialMapping
from app.models.material import Material
from app.models.cpse import CPSE
from app.models.inventory import Inventory
from app.models.audit_log import AuditLog
from app.models.material_cluster import MaterialCluster
from app.models.material_request import MaterialRequest
from app.models.user import User
from app.services.audit_service import verify_audit_chain

def rows(db):
    mappings_by_national = defaultdict(list)
    for mapping in db.query(MaterialMapping).all():
        mappings_by_national[mapping.national_material_id].append(mapping)
    materials = {material.id: material for material in db.query(Material).all()}
    cpses = {cpse.id: cpse for cpse in db.query(CPSE).all()}
    quantity_by_material = defaultdict(float)
    for inventory in db.query(Inventory).all():
        quantity_by_material[inventory.material_id] += max(
            inventory.available_quantity - inventory.reserved_quantity,
            0,
        )

    output = []
    for n in db.query(NationalMaterial).all():
        mappings = mappings_by_national[n.id]

        if not mappings:
            output.append({
                "National Code": n.national_code,
                "Standard Description": n.description,
                "Category": n.category,
                "UOM": n.unit,
                "CPSE": None,
                "Legacy Code": None,
                "Available Quantity": 0,
                "Mapping Type": None,
            })

        for mp in mappings:
            material = materials.get(mp.material_id)
            cpse = cpses.get(material.cpse_id) if material else None
            qty = quantity_by_material[mp.material_id]
            output.append({
                "National Code": n.national_code,
                "Standard Description": n.description,
                "Category": n.category,
                "UOM": n.unit,
                "CPSE": cpse.name if cpse else None,
                "Legacy Code": material.material_code if material else None,
                "Available Quantity": qty,
                "Mapping Type": mp.mapping_type,
            })
    return output

def export_csv(db):
    s = StringIO()
    pd.DataFrame(rows(db)).to_csv(s, index=False)
    return s.getvalue().encode()

def export_xlsx(db):
    b = BytesIO()
    with pd.ExcelWriter(b, engine="openpyxl") as writer:
        pd.DataFrame(rows(db)).to_excel(writer, index=False, sheet_name="National Registry")
    return b.getvalue()


def _text(value, limit=180):
    if value is None:
        return "—"
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str)
    value = str(value).replace("\n", " ").strip()
    return value if len(value) <= limit else f"{value[:limit - 1]}…"


def _timestamp(value):
    return value.strftime("%d %b %Y, %H:%M UTC") if value else "—"


def governance_report_sections(db):
    """Authoritative, table-oriented snapshot for controlled PDF/DOCX exports."""
    cpses = {row.id: row for row in db.query(CPSE).all()}
    users = {row.id: row for row in db.query(User).all()}
    materials = {row.id: row for row in db.query(Material).all()}
    national_materials = {row.id: row for row in db.query(NationalMaterial).all()}

    audit_rows = []
    for row in db.query(AuditLog).order_by(AuditLog.id.desc()).all():
        audit_rows.append([
            row.id, _timestamp(row.created_at), row.action, row.entity_type,
            row.entity_id or "—", users.get(row.user_id).email if row.user_id in users else "System",
            _text(row.details), row.current_hash or "Legacy / unhashed",
        ])

    account_rows = [[
        row.id, row.name, row.email, row.account_status, row.role,
        cpses.get(row.cpse_id).code if row.cpse_id in cpses else "—",
        _timestamp(row.created_at), _timestamp(row.reviewed_at),
    ] for row in sorted(users.values(), key=lambda item: item.id)]

    registry_rows = [[
        row.national_code or "—", row.description, row.category or "—", row.unit or "—",
        row.status, _timestamp(row.updated_at),
    ] for row in db.query(NationalMaterial).order_by(NationalMaterial.national_code).all()]

    inventory_rows = [[
        cpses.get(row.cpse_id).code if row.cpse_id in cpses else "—",
        materials.get(row.material_id).material_code if row.material_id in materials else f"Material #{row.material_id}",
        row.warehouse, row.available_quantity, row.reserved_quantity,
        max(row.available_quantity - row.reserved_quantity, 0), row.uom, _timestamp(row.last_updated),
    ] for row in db.query(Inventory).order_by(Inventory.id).all()]

    request_rows = [[
        row.request_number or f"Request #{row.id}", row.status,
        cpses.get(row.requesting_cpse_id).code if row.requesting_cpse_id in cpses else "—",
        users.get(row.requested_by).email if row.requested_by in users else "—",
        _timestamp(row.created_at), _timestamp(row.submitted_at), _timestamp(row.approved_at),
    ] for row in db.query(MaterialRequest).order_by(MaterialRequest.id.desc()).all()]

    cluster_rows = [[
        row.cluster_code or f"Cluster #{row.id}", row.status, row.risk_level,
        national_materials.get(row.national_material_id).national_code if row.national_material_id in national_materials else "—",
        _timestamp(row.created_at), _timestamp(row.approved_at), _text(row.approval_comment or row.review_comment, 100),
    ] for row in db.query(MaterialCluster).order_by(MaterialCluster.id.desc()).all()]

    integrity = verify_audit_chain(db)
    summary = [
        ["Generated at", _timestamp(datetime.utcnow())],
        ["Audit ledger integrity", integrity.get("status", "UNKNOWN")],
        ["Audit events", len(audit_rows)],
        ["User accounts", len(account_rows)],
        ["National materials", len(registry_rows)],
        ["Inventory records", len(inventory_rows)],
        ["Material requests", len(request_rows)],
        ["Identity clusters", len(cluster_rows)],
    ]
    return [
        ("Report summary", ["Metric", "Value"], summary),
        ("Audit & governance ledger", ["ID", "Timestamp", "Action", "Entity", "Entity ID", "Actor", "Details", "Integrity hash"], audit_rows),
        ("User accounts", ["ID", "Name", "Email", "Account status", "Role", "CPSE", "Registered", "Reviewed"], account_rows),
        ("National material registry", ["National code", "Description", "Category", "UOM", "Status", "Updated"], registry_rows),
        ("Cross-CPSE inventory", ["CPSE", "Material code", "Warehouse", "Available", "Reserved", "Effective", "UOM", "Updated"], inventory_rows),
        ("Material requests", ["Request", "Status", "CPSE", "Requested by", "Created", "Submitted", "Approved"], request_rows),
        ("Identity cluster governance", ["Cluster", "Status", "Risk", "National code", "Created", "Approved", "Review note"], cluster_rows),
    ]


def export_governance_pdf(db):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=landscape(A4), rightMargin=10 * mm, leftMargin=10 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#20354a"), fontSize=18, leading=22)
    heading = ParagraphStyle("ReportHeading", parent=styles["Heading2"], textColor=colors.HexColor("#20354a"), spaceBefore=10, spaceAfter=6)
    cell = ParagraphStyle("ReportCell", parent=styles["BodyText"], fontSize=6.5, leading=8)
    story = [Paragraph("National Unified Material Master", title), Paragraph("Governance, Audit and Operational Data Report", styles["Heading2"]), Spacer(1, 4)]
    for index, (section, headers, rows) in enumerate(governance_report_sections(db)):
        story.append(Paragraph(section, heading))
        table_data = [[Paragraph(_text(value, 90), cell) for value in headers]]
        table_data.extend([[Paragraph(_text(value), cell) for value in row] for row in rows] or [[Paragraph("No records available", cell) for _ in headers]])
        table = Table(table_data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#20354a")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#C9D2DA")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        story.append(table)
        if index < 6:
            story.append(PageBreak())
    document.build(story)
    return stream.getvalue()


def export_governance_docx(db):
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.left_margin = section.right_margin = Inches(0.45)
    title = document.add_heading("National Unified Material Master", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Governance, Audit and Operational Data Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for section_name, headers, rows in governance_report_sections(db):
        document.add_heading(section_name, level=1)
        table = document.add_table(rows=1, cols=len(headers))
        table.style = "Light Shading Accent 1"
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = _text(value, 90)
        for row in rows or [["No records available"] + [""] * (len(headers) - 1)]:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = _text(value)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(7)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
